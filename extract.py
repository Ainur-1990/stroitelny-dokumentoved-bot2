"""Извлечение текста из документов, присланных в Telegram.

Поддерживаются PDF и DOCX. Фото обрабатываются отдельным vision-путём в app.py.
"""
import io

from pypdf import PdfReader
from pypdf.errors import PdfReadError
from docx import Document


MAX_FILE_BYTES = 20 * 1024 * 1024  # лимит Bot API на скачивание файла
MAX_TEXT_CHARS = 30000
MIN_MEANINGFUL_CHARS = 20

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class ExtractionError(Exception):
    """Ошибка разбора; текст безопасен для показа пользователю."""


def _pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractionError("PDF защищён паролем. Снимите защиту и пришлите файл заново.")
        pages = [(page.extract_text() or "") for page in reader.pages]
    except ExtractionError:
        raise
    except (PdfReadError, ValueError, OSError):
        raise ExtractionError("Не удалось прочитать PDF: файл повреждён или в необычном формате.")
    text = "\n".join(pages).strip()
    if len(text) < MIN_MEANINGFUL_CHARS:
        raise ExtractionError(
            "В PDF нет текстового слоя (похоже, это скан). Сфотографируйте документ и пришлите фото — "
            "оно уйдёт на распознавание в ИИ."
        )
    return text


def _docx(data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(data))
    except (ValueError, KeyError, OSError):
        raise ExtractionError("Не удалось прочитать DOCX: файл повреждён или в старом формате .doc.")
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    text = "\n".join(part for part in parts if part.strip()).strip()
    if len(text) < MIN_MEANINGFUL_CHARS:
        raise ExtractionError("В DOCX не нашлось текста для анализа.")
    return text


def extract_text(data: bytes, mime: str, filename: str) -> str:
    """Возвращает текст документа или бросает ExtractionError с причиной."""
    name = (filename or "").lower()
    if mime == "application/pdf" or name.endswith(".pdf"):
        text = _pdf(data)
    elif mime == DOCX_MIME or name.endswith((".docx", ".docm")):
        text = _docx(data)
    else:
        raise ExtractionError(
            "Поддерживаются PDF, DOCX и фото документов. Другие форматы (xls, dwg, zip) пока не обрабатываются — "
            "опишите задачу текстом."
        )
    return text[:MAX_TEXT_CHARS]
